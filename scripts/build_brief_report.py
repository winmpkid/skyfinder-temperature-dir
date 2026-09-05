from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "reports"
DOCX_PATH = OUTPUT / "SkyFinder_DIR_Brief_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_borders(paragraph):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "between"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead:
        paragraph.add_run(bold_lead).bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def build_document():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)

    for style_name, size in (("Heading 1", 13.5), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)

    title_paragraph = doc.add_paragraph(
        "SkyFinder Temperature Prediction with Deep Imbalanced Regression",
        style="Title",
    )
    remove_paragraph_borders(title_paragraph)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("Brief take-home experiment report for the Health Intelligence Lab")
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(70, 70, 70)

    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(7)
    summary.add_run("Main finding  ").bold = True
    summary.add_run(
        "Label Distribution Smoothing reduced validation MAE from 2.677 °C to "
        "2.610 °C, but test MAE increased from 4.922 °C to 6.315 °C. The result "
        "suggests that label imbalance matters in some temperature ranges, while "
        "temporal distribution shift is the larger generalization problem in this subset."
    )

    doc.add_heading("Objective", level=1)
    add_body(
        doc,
        "I adapted the Label Distribution Smoothing component of Deep Imbalanced "
        "Regression to predict ambient temperature from SkyFinder images. The "
        "experiment compares a standard ResNet-18 regressor with the same model "
        "trained using LDS-weighted loss."
    )

    doc.add_heading("Method", level=1)
    add_body(
        doc,
        "Images from cameras 858, 3888, and 4795 were joined with the SkyFinder "
        "weather metadata. "
        "Invalid temperatures and two truncated JPEG files were removed. Samples "
        "were split chronologically by capture date into 1,777 training, 375 "
        "validation, and 395 test images, which prevents adjacent frames from the "
        "same date from being divided across splits."
    )
    add_body(
        doc,
        "The backbone was an ImageNet-pretrained ResNet-18 with its classification "
        "layer replaced by a one-output regression head. Images were resized to "
        "224 by 224 pixels and normalized with ImageNet statistics; random "
        "horizontal flipping was applied only during training. Both models used "
        "AdamW, a learning rate of 0.0001, weight decay of 0.0001, batch size 32, "
        "20 epochs, and random seed 42. The baseline minimized mean absolute error."
    )
    add_body(
        doc,
        "For LDS, training temperatures were placed into 1 °C bins. Their counts "
        "were smoothed with a Gaussian kernel of size 5 and sigma 2. Each sample "
        "received a weight inversely proportional to the smoothed density of its "
        "temperature bin; weights were normalized to mean 1. The weighted absolute "
        "error was used only for optimization, while ordinary MAE was retained for "
        "validation and checkpoint selection."
    )

    results_heading = doc.add_heading("Experimental Results", level=1)
    results_heading.paragraph_format.page_break_before = True
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.05), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)]
    headers = ["Model", "Val MAE", "Val RMSE", "Test MAE", "Test RMSE"]
    for index, (cell, label, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.2)

    rows = [
        ("Training median", "10.226", "11.010", "5.945", "7.803"),
        ("ResNet-18 baseline", "2.677", "3.459", "4.922", "6.293"),
        ("ResNet-18 with LDS", "2.610", "3.312", "6.315", "7.551"),
    ]
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, "EEF4F8")
        for index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = width
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.size = Pt(9.2)
        prevent_row_split(table.rows[-1])
    set_table_borders(table)
    caption = doc.add_paragraph("Table 1  Temperature error in degrees Celsius")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(4)
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8.8)

    add_body(
        doc,
        "On validation data, LDS improved MAE by 0.066 °C (2.5 percent) and RMSE "
        "by 0.148 °C (4.3 percent). Its largest validation gains were below 0 °C: "
        "MAE fell from 4.175 °C to 3.558 °C in the -10 to -5 °C bin and from "
        "3.409 °C to 2.432 °C in the -5 to 0 °C bin."
    )

    doc.add_heading("Analysis", level=1)
    add_body(
        doc,
        "The validation result supports the intended role of LDS: neighboring "
        "temperature labels share information, and rare low-temperature examples "
        "receive greater influence during training. The improvement was not uniform. "
        "For example, validation MAE in the 5 to 10 °C bin rose from 6.316 °C to "
        "7.706 °C, showing that inverse-density weighting can trade performance "
        "between regions rather than improve the complete range."
    )
    add_body(
        doc,
        "The chronological test split was substantially colder than the training "
        "and validation splits: mean temperatures were 5.89 °C, 11.36 °C, and "
        "13.25 °C, respectively. Both models overpredicted on the test set, but the "
        "mean prediction bias was larger with LDS (+4.95 °C) than with the baseline "
        "(+3.82 °C). Consequently, LDS failed to generalize despite its validation "
        "gain. This indicates that temporal covariate shift and limited data coverage "
        "were more important than label frequency alone in the final split."
    )

    doc.add_heading("Improvements", level=1)
    add_bullet(
        doc,
        "Tune LDS bin width, kernel size, sigma, and a maximum weight cap on the validation set; the current maximum sample weight was about 24.1."
    )
    add_bullet(
        doc,
        "Implement Feature Distribution Smoothing and compare baseline, LDS, FDS, and LDS plus FDS under the same protocol."
    )
    add_bullet(
        doc,
        "Add more cameras and seasons, or design seasonal splits, to reduce the train-test temperature shift while preserving temporal separation."
    )
    add_bullet(
        doc,
        "Run multiple random seeds and report uncertainty; examine time of day, weather metadata, and prediction calibration as possible sources of error."
    )

    doc.add_heading("References", level=1)
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(2)
    p1.add_run(
        "1. Yang, Y., Zha, K., Chen, Y., Wang, H., and Katabi, D. "
        "Delving into Deep Imbalanced Regression. ICML, 2021. "
    )
    add_hyperlink(p1, "PMLR paper", "https://proceedings.mlr.press/v139/yang21m.html")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run("2. Mihail, R. P. et al. SkyFinder dataset. ")
    add_hyperlink(p2, "Dataset page", "https://mvrl.cse.wustl.edu/datasets/skyfinder/")

    core = doc.core_properties
    core.title = "SkyFinder Temperature Prediction with Deep Imbalanced Regression"
    core.subject = "Health Intelligence Lab take-home experiment report"
    core.keywords = "SkyFinder, temperature regression, ResNet-18, LDS, DIR"
    core.author = ""
    core.last_modified_by = ""

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
